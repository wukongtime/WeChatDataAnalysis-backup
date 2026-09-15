"""跳过图片前不读取或转换图片，保留附件文字及实际模型的缓存隔离。"""
import asyncio
import io
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from PIL import Image

from wechat_decrypt_tool.ai import media
from wechat_decrypt_tool.ai.storage import AIStore


def png():
    stream = io.BytesIO()
    Image.new('RGB', (4, 4), 'red').save(stream, 'PNG')
    return stream.getvalue()


def message(kind='file'):
    return {'kind': kind, 'text': '[附件]', 'username': 'friend', 'anchor': 'one', 'media': {}}


def test_unsupported_image_never_reads_or_converts(tmp_path, monkeypatch):
    monkeypatch.setattr(media, 'resolve_media', lambda *args: pytest.fail('不应读取图片'))
    monkeypatch.setattr(media, 'image_url', lambda *args: pytest.fail('不应转换图片'))
    models = SimpleNamespace(invoke=AsyncMock())
    service = media.MediaService(AIStore(tmp_path), models)
    result = asyncio.run(service.enrich('account', message('image'), {'skip_unsupported_images': True}, {}, lambda: None))
    assert result['coverage'] == '当前模型不支持图片，已跳过图片内容'
    assert result['text'] == '[附件]'
    models.invoke.assert_not_called()


@pytest.mark.parametrize('suffix', ['.docx', '.pptx', '.xlsx', '.pdf'])
def test_document_skip_preserves_text_without_image_conversion(tmp_path, monkeypatch, suffix):
    stream = io.BytesIO()
    if suffix == '.docx':
        from docx import Document
        doc = Document(); doc.add_paragraph('text-before'); doc.add_picture(io.BytesIO(png())); doc.add_paragraph('text-after'); doc.save(stream)
    elif suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = 'text-before'
        slide.shapes.add_picture(io.BytesIO(png()), 0, 0)
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = 'text-after'; deck.save(stream)
    elif suffix == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage
        book = Workbook(); book.active['A1'] = 'text-before'; book.active['A2'] = 'text-after'
        book.active.add_image(SheetImage(io.BytesIO(png())), 'B1'); book.save(stream)
    else:
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.write(stream)
        import pypdfium2
        monkeypatch.setattr(pypdfium2, 'PdfDocument', lambda *args: pytest.fail('不应渲染扫描页'))
    monkeypatch.setattr(media, 'resolve_media', lambda *args: (stream.getvalue(), suffix))
    original_open = zipfile.ZipFile.open
    def open_without_images(archive, name, *args, **kwargs):
        filename = name.filename if isinstance(name, zipfile.ZipInfo) else name
        assert '/media/' not in filename, '纯文字解析不应解压图片部件'
        return original_open(archive, name, *args, **kwargs)
    monkeypatch.setattr(zipfile.ZipFile, 'open', open_without_images)
    monkeypatch.setattr(media, 'image_url', lambda *args: pytest.fail('不应转换内嵌图片'))
    models = SimpleNamespace(invoke=AsyncMock())
    service = media.MediaService(AIStore(tmp_path), models)
    for _ in range(2):
        result = asyncio.run(service.enrich('account', message(), {'skip_unsupported_images': True}, {}, lambda: None))
        assert '已跳过图片内容' in result['coverage']
        if suffix != '.pdf':
            assert 'text-before' in result['text'] and 'text-after' in result['text']
    models.invoke.assert_not_called()


def test_image_cache_uses_actual_model_id(tmp_path, monkeypatch):
    monkeypatch.setattr(media, 'resolve_media', lambda *args: (png(), '.image'))
    models = SimpleNamespace(invoke=AsyncMock(side_effect=['from-a', 'from-b']))
    service = media.MediaService(AIStore(tmp_path), models)
    async def check():
        for name, expected in [('a', 'from-a'), ('b', 'from-b'), ('a', 'from-a')]:
            result = await service.enrich('account', message('image'), {'skip_unsupported_images': True},
                {'id': 'service', 'model': name, 'vision': True, 'revision': 1}, lambda: None)
            assert expected in result['text']
        assert [call.args[0]['model'] for call in models.invoke.call_args_list] == ['a', 'b']
    asyncio.run(check())
