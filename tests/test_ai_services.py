import asyncio
import io
import json
import sys
import time
from pathlib import Path
from unittest.mock import patch

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from wechat_decrypt_tool.ai.media import parse_document, image_url
from wechat_decrypt_tool.ai.messages import filter_after, advance_cursor
from wechat_decrypt_tool.ai.providers import ModelService, ProviderFailure, public_profile, validate_url
from wechat_decrypt_tool.ai.schemas import TaskInput, RuleInput
from wechat_decrypt_tool.ai.service import AIService, split_messages, validate_sources, next_due
from wechat_decrypt_tool.ai.storage import AIStore


def message(id="a", timestamp=100, text="项目延期，需要确认明天下午交付"):
    return {"source": id, "identity": id, "time": timestamp, "username": "group", "anchor": "db:table:1", "sender": "张三", "text": text, "kind": "text", "media": {}}


class FakeModels(ModelService):
    def __init__(self, store):
        super().__init__(store)
        self.calls = []
        self.fail = False

    async def invoke(self, profile, prompt, schema=None, images=None, account=""):
        self.calls.append(prompt)
        if self.fail:
            raise ProviderFailure("模拟连接失败")
        if schema and schema.__name__ == "Matches":
            return {"matches": [{"reason": "讨论交付延期", "sources": ["a"]}]}
        return {"overview": "项目交付需要确认", "topics": [{"text": "交付时间待确认", "sources": ["a"]}], "conclusions": [], "todos": []}


@pytest.fixture
def service(tmp_path):
    store = AIStore(tmp_path)
    store.put("profile", {"name": "测试", "provider": "custom", "protocol": "openai", "base_url": "http://localhost:1234/v1", "api_key": "secret-value", "model": "fake", "vision": False, "revision": 1}, id="model")
    store.put("defaults", {"text": "model", "vision": ""}, id="global")
    def reader(account, username, start, end, count=None, require_realtime=False, checkpoint=None):
        values = [message()] if (start is None or start <= 100) and end >= 100 else []
        if checkpoint:
            checkpoint()
        return {"username": username, "name": "项目群", "messages": values, "warning": ""}
    return AIService(store, FakeModels(store), reader)


def task_options(**kwargs):
    return {"account": "account", "conversations": ["group"], "media": False, **kwargs}


def test_restart_closes_orphan_usage_without_fabricating_usage_or_duration(tmp_path):
    store = AIStore(tmp_path)
    for account, usage in (("first", {}), ("second", {"input_tokens": 12})):
        store.put("usage", {"status": "running", "started_at": 10, "usage": usage,
                            "usage_known": bool(usage)}, id=account, account=account)
    finished = store.put("usage", {"status": "success", "finished_at": 20,
                                  "duration_ms": 5}, id="finished", account="first")
    store.recover_interrupted_usage()
    recovered = store.list("usage")
    for account in ("first", "second"):
        row = store.get("usage", account)
        assert row["status"] == "interrupted"
        assert row["error_type"] == "ProcessInterrupted"
        assert "finished_at" not in row and "duration_ms" not in row
    assert store.get("usage", "first")["usage"] == {}
    assert store.get("usage", "first")["usage_known"] is False
    assert store.get("usage", "second")["usage"] == {"input_tokens": 12}
    assert store.get("usage", "second")["usage_known"] is True
    assert store.get("usage", "finished") == finished
    store.recover_interrupted_usage()
    assert store.list("usage") == recovered


def test_alert_context_respects_recent_window_and_explicit_start(service):
    now = 2_000_000
    calls = []
    def reader(account, username, start, end, count=None, require_realtime=False, checkpoint=None):
        calls.append((start, end, count, require_realtime))
        candidates = [message("old", now - 86401), message("recent", now - 100), message("new", now)]
        return {"username": username, "messages": [m for m in candidates if start <= m["time"] <= end]}
    service.reader = reader
    rule = {"kind": "alert", "trigger": "interval", "cursors": {"group": {"time": now - 1, "ids": []}}}
    options = {**task_options(), "range": {"mode": "count", "count": 100, "end": now, "start": None}}
    result = asyncio.run(service.read_conversations(options, rule))
    assert [m["source"] for m in result[0]["messages"]] == ["new"]
    assert [m["source"] for m in result[0]["context"]] == ["recent"]
    assert calls[-1] == (now - 86400, now, 20, True)
    options["range"]["start"] = now - 50
    result = asyncio.run(service.read_conversations(options, rule))
    assert result[0]["context"] == []
    assert calls[-1][0] == now - 50


def test_range_validation_and_same_second_cursor():
    with pytest.raises(ValueError):
        TaskInput.model_validate(task_options(range={"mode": "dates", "start": 20, "end": 10}))
    values = [message("old", 100), message("same-second", 100), message("new", 101)]
    cursor = {"time": 100, "ids": ["old"]}
    assert [m["source"] for m in filter_after(values, cursor)] == ["same-second", "new"]
    assert advance_cursor(values, cursor) == {"time": 101, "ids": ["new"]}


def test_split_long_attachment_preserves_entire_text():
    text = "项目资料" * 15000
    chunks = split_messages([message(text=text)])
    assert len(chunks) > 1
    assert "".join(m["text"] for c in chunks for m in c) == text
    assert all(m["source"] == "a" for c in chunks for m in c)


def test_reject_invented_sources():
    with pytest.raises(ProviderFailure):
        validate_sources({"matches": [{"reason": "x", "sources": ["invented"]}]}, {"a"}, True)


def test_profile_mask_and_url_validation(service):
    assert "api_key" not in public_profile(service.models.resolve())
    assert service.models.resolve()["api_key"] == "secret-value"
    for value in ("file:///etc/passwd", "https://user:key@example.com", "http://example.com", "https://example.com/?key=secret"):
        with pytest.raises(ValueError):
            validate_url(value)
    validate_url("http://localhost:11434/v1")


def test_summary_graph_and_checkpoint_replay(service):
    async def run():
        task = service.create_task(task_options())
        assert "secret-value" not in json.dumps(task)
        await service.execute(task["id"])
        saved = service.store.get("task", task["id"])
        assert saved["status"] == "completed", saved
        assert saved["results"][0]["count"] == 1
        calls = len(service.models.calls)
        # 模拟工作流已完成、业务投递阶段中断后重新执行。
        service.update_task(task["id"], status="running")
        await service.execute(task["id"])
        assert len(service.models.calls) == calls
        assert len(service.store.events(pending=True)) == 1
        checkpoint_bytes = (service.store.root / "checkpoints.sqlite3").read_bytes()
        assert b"secret-value" not in checkpoint_bytes
    asyncio.run(run())


def test_live_stages_visible_before_model_returns_and_finish_time(service):
    async def run():
        image_entered, image_release = asyncio.Event(), asyncio.Event()
        model_entered, model_release = asyncio.Event(), asyncio.Event()
        original = service.models.invoke
        async def enrich(account, value, *args):
            image_entered.set()
            await image_release.wait()
            return value
        async def invoke(*args, **kwargs):
            model_entered.set()
            await model_release.wait()
            return await original(*args, **kwargs)
        service.media.enrich = enrich
        service.models.invoke = invoke
        task = service.create_task(task_options(media=True), prepared=[{
            "username": "group", "name": "项目群", "messages": [{**message(), "kind": "image"}], "context": []}])
        worker = asyncio.create_task(service.execute(task["id"]))
        try:
            await asyncio.wait_for(image_entered.wait(), 5)
            image_task = service.store.get("task", task["id"])
            assert image_task["stage"] == "处理图片与附件 1/1 · 项目群"
            assert 0 < image_task["progress"] < 100
            assert image_task["message_count"] == 1
            image_release.set()
            await asyncio.wait_for(model_entered.wait(), 5)
            model_task = service.store.get("task", task["id"])
            assert model_task["stage"] == "生成摘要 1/1 · 项目群"
            assert image_task["progress"] < model_task["progress"] < 100
            assert model_task["finished_at"] is None
            model_release.set()
            await asyncio.wait_for(worker, 5)
            finished = service.store.get("task", task["id"])
            assert finished["status"] == "completed"
            assert finished["finished_at"] >= finished["started_at"]
            assert finished["activity"][-1]["stage"] == "处理完成"
            assert "secret-value" not in json.dumps(finished["activity"])
        finally:
            worker.cancel()
            await asyncio.gather(worker, return_exceptions=True)
    asyncio.run(run())


def test_failure_does_not_advance_rule_cursor_and_retry(service):
    async def run():
        rule = await service.save_rule({**task_options(), "name": "项目提醒", "kind": "alert", "condition": "交付延期"})
        rule["cursors"] = {"group": {"time": 1, "ids": []}}
        service.store.put("rule", rule, id=rule["id"])
        task = await service.run_rule(rule, force=True)
        service.models.fail = True
        await service.execute(task["id"])
        assert service.store.get("rule", rule["id"])["cursors"]["group"]["time"] == 1
        assert service.store.get("task", task["id"])["status"] == "failed"
        service.models.fail = False
        await service.reset_graph(task["id"])
        await service.execute(task["id"])
        assert service.store.get("rule", rule["id"])["cursors"]["group"]["time"] == 100
        assert len(service.store.list("alert", "account")) == 1
    asyncio.run(run())


def test_no_new_messages_no_call_and_threshold(service):
    async def run():
        rule = await service.save_rule({**task_options(), "name": "阈值", "trigger": "count", "threshold": 2})
        rule["cursors"] = {"group": {"time": 1, "ids": []}}
        assert await service.run_rule(rule) is None
        rule["threshold"] = 1
        assert await service.run_rule(rule) is not None
        assert service.models.calls == []
    asyncio.run(run())


def test_account_isolation_and_cleanup(service):
    task = service.create_task(task_options())
    service.store.put("alert", {"reason": "other"}, account="other")
    service.purge_account("account")
    assert service.store.get("task", task["id"]) is None
    assert len(service.store.list("alert", "other")) == 1


def test_docx_xlsx_pptx_and_pdf_parsers():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter
    document = Document(); document.add_paragraph("中文项目说明")
    table = document.add_table(rows=1, cols=1); table.cell(0, 0).text = "交付时间"
    data = io.BytesIO(); document.save(data)
    assert "交付时间" in json.dumps(parse_document(data.getvalue(), ".docx"), ensure_ascii=False)
    book = Workbook(); book.active.title = "项目"; book.active.append(["任务", "截止日"])
    data = io.BytesIO(); book.save(data)
    assert "工作表 项目" in json.dumps(parse_document(data.getvalue(), ".xlsx"), ensure_ascii=False)
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0]); slide.shapes.title.text = "进展报告"
    data = io.BytesIO(); deck.save(data)
    assert "进展报告" in json.dumps(parse_document(data.getvalue(), ".pptx"), ensure_ascii=False)
    pdf = PdfWriter(); pdf.add_blank_page(width=100, height=100)
    data = io.BytesIO(); pdf.write(data)
    assert any("image" in part for part in parse_document(data.getvalue(), ".pdf"))


def test_invalid_and_encrypted_documents():
    from pypdf import PdfWriter
    pdf = PdfWriter(); pdf.add_blank_page(width=100, height=100); pdf.encrypt("secret")
    data = io.BytesIO(); pdf.write(data)
    with pytest.raises(ValueError, match="加密"):
        parse_document(data.getvalue(), ".pdf")
    with pytest.raises(ValueError, match="不支持"):
        parse_document(b"", ".doc")


def test_gif_uses_first_frame():
    from PIL import Image
    first, second = Image.new("RGB", (20, 20), "red"), Image.new("RGB", (20, 20), "blue")
    data = io.BytesIO(); first.save(data, "GIF", save_all=True, append_images=[second])
    import base64
    image = Image.open(io.BytesIO(base64.b64decode(image_url(data.getvalue()).split(",")[1])))
    r, g, b = image.getpixel((0, 0))
    assert r > b


def test_cancel_running_model_request(service):
    async def run():
        entered = asyncio.Event()
        async def waiting(*args, **kwargs):
            entered.set()
            await asyncio.Event().wait()
        service.models.invoke = waiting
        task = service.create_task(task_options())
        service.worker = asyncio.create_task(service.execute(task["id"]))
        await asyncio.wait_for(entered.wait(), timeout=5)
        await service.cancel_task(task["id"])
        assert service.worker.done()
        assert service.store.get("task", task["id"])["status"] == "cancelled"
        assert service.store.events(pending=True) == []
    asyncio.run(run())


def test_concurrent_rule_dispatch_creates_single_task(service):
    async def run():
        rule = await service.save_rule({**task_options(), "name": "项目", "kind": "alert", "condition": "延期"})
        rule["cursors"] = {"group": {"time": 1, "ids": []}}
        service.store.put("rule", rule, id=rule["id"])
        results = await asyncio.gather(service.run_rule(rule), service.run_rule(rule))
        assert len([r for r in results if r]) == 1
    asyncio.run(run())


def test_recent_count_uses_unique_rows_and_bounded_windows(tmp_path):
    from wechat_decrypt_tool.ai.messages import read_messages
    from wechat_decrypt_tool import chat_helpers, chat_export_service, account_source_policy
    Row = chat_export_service._Row
    rows = [Row("db", "table", i, i, 1, i, 100 + i, f"消息{i}", "sender", False) for i in range(1, 201)]
    rows.append(rows[-1])
    windows = []
    def read(**kwargs):
        windows.append((kwargs["start_time"], kwargs["end_time"]))
        return iter(r for r in rows if (kwargs["start_time"] is None or r.create_time >= kwargs["start_time"]) and r.create_time <= kwargs["end_time"])
    with patch.object(chat_helpers, "_resolve_account_dir", return_value=tmp_path), patch.object(account_source_policy, "account_prefers_decrypted_snapshot", return_value=True), patch.object(chat_export_service, "_iter_rows_for_conversation", side_effect=read):
        result = read_messages("account", "group", None, 500, 100)
    assert len(result["messages"]) == 100
    assert {m["time"] for m in result["messages"]} == set(range(201, 301))
    assert len(windows) == 1


def test_restart_does_not_reset_monitoring_baseline(service):
    async def run():
        rule = await service.save_rule({**task_options(), "name": "项目", "kind": "alert", "condition": "延期"})
        restored = AIService(AIStore(service.store.root), service.models, service.reader)
        assert restored.store.get("rule", rule["id"])["cursors"] == rule["cursors"]
    asyncio.run(run())


def test_failed_automatic_task_backoff_and_pause(service):
    async def run():
        rule = await service.save_rule({**task_options(), "name": "项目", "kind": "alert", "condition": "延期"})
        rule["enabled"] = True; rule["cursors"] = {"group": {"time": 1, "ids": []}}
        service.store.put("rule", rule, id=rule["id"])
        task = await service.run_rule(rule)
        service.models.fail = True
        for i in range(3):
            await service.reset_graph(task["id"])
            await service.execute(task["id"])
            saved = service.store.get("task", task["id"])
            assert saved["attempts"] == i + 1
            if i < 2:
                assert saved["retry_at"] > time.time()
        assert service.store.get("rule", rule["id"])["enabled"] is False
    asyncio.run(run())
