"""离线、合成数据运行检查：不读取账号，不连接模型服务。"""
import asyncio
import io
import platform
import sys
import tempfile
from pathlib import Path


async def _checkpoint(root):
    from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
    from langgraph.graph import StateGraph, START, END
    from typing import TypedDict

    class State(TypedDict):
        count: int

    async with AsyncSqliteSaver.from_conn_string(str(root / 'checkpoint.sqlite3')) as saver:
        graph = StateGraph(State)
        graph.add_node('increment', lambda state: {'count': state['count'] + 1})
        graph.add_edge(START, 'increment')
        graph.add_edge('increment', END)
        workflow = graph.compile(checkpointer=saver)
        config = {'configurable': {'thread_id': 'synthetic-runtime-check'}}
        assert (await workflow.ainvoke({'count': 0}, config))['count'] == 1
        assert (await workflow.aget_state(config)).values['count'] == 1


def check_runtime(model_root=None):
    from PIL import Image
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from langchain_openai import ChatOpenAI
    from langchain_anthropic import ChatAnthropic
    from langchain_core.messages import HumanMessage
    import huggingface_hub
    import onnxruntime
    import tokenizers
    import tiktoken
    from .media import parse_document, image_url
    from .storage import AIStore
    from ..local_search.catalog import model_spec, model_dir
    from ..local_search.index import SemanticIndex
    from ..local_search.inference import LocalInference

    report = {'ok': False, 'platform': sys.platform, 'arch': platform.machine(),
              'frozen': bool(getattr(sys, 'frozen', False)), 'onnxruntime': onnxruntime.__version__}
    # 仅构造适配器和工具定义，不发起任何线上请求。
    tool = {'name': 'chat_action', 'description': '合成检查', 'parameters': {'type': 'object', 'properties': {}}}
    for cls in (ChatOpenAI, ChatAnthropic):
        cls(model='synthetic-test-model', api_key='synthetic-unused-key').bind_tools([tool])
    assert HumanMessage('检查').content == '检查'
    assert 'cl100k_base' in tiktoken.list_encoding_names()
    with tempfile.TemporaryDirectory(prefix='wechat-ai-runtime-') as directory:
        root = Path(directory)
        store = AIStore(root / 'business')
        store.put('runtime_check', {'ok': True}, id='check', account='synthetic')
        assert store.get('runtime_check', 'check')['ok']
        asyncio.run(_checkpoint(root))
        index = SemanticIndex(root / 'vectors.sqlite3')
        with index.connection() as db:
            assert db.execute("SELECT vec_distance_cosine('[1,0]','[1,0]')").fetchone()[0] == 0
            report['sqlite_vec'] = db.execute('SELECT vec_version()').fetchone()[0]
            db.execute('CREATE VIRTUAL TABLE runtime_fts USING fts5(text)')
        document = Document(); document.add_paragraph('合成文档')
        workbook = Workbook(); workbook.active.append(['合成表格', 42])
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = '合成幻灯片'
        parsed = []
        for suffix, item in [('.docx', document), ('.xlsx', workbook), ('.pptx', deck)]:
            output = io.BytesIO(); item.save(output)
            assert any(part.get('text') for part in parse_document(output.getvalue(), suffix))
            parsed.append(suffix)
        picture = Image.new('RGB', (64, 64), 'green')
        for format in ['JPEG', 'PNG', 'WEBP', 'GIF']:
            output = io.BytesIO(); picture.save(output, format)
            assert image_url(output.getvalue()).startswith('data:image/jpeg;')
        output = io.BytesIO(); picture.save(output, 'PDF')
        assert any(part.get('image') for part in parse_document(output.getvalue(), '.pdf'))
        report['media'] = parsed + ['.pdf', 'JPEG', 'PNG', 'WEBP', 'GIF']
        report['catalog'] = [model_spec(name)['id'] for name in ['bge-small-zh', 'bge-base-zh', 'e5-small']]
        if model_root:
            engine = LocalInference()
            try:
                spec = model_spec('bge-small-zh')
                vectors = engine.encode(model_dir(Path(model_root), spec['id']), spec, ['本地报价查找'], strategy='cpu')
                assert len(vectors[0]) == spec['dimension']
                report['inference'] = {**engine.status, 'dimension': len(vectors[0])}
            finally:
                engine.close()
    report['ok'] = True
    return report
