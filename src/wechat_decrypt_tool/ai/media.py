from __future__ import annotations
from .diagnostics import observed, event as diagnostic_event
import logging

import asyncio
import base64
import hashlib
import io
import json
import time
import zipfile
from pathlib import Path

from PIL import Image


def image_url(data):
    with Image.open(io.BytesIO(data)) as img:
        img.seek(0)
        img = img.convert("RGB")
        img.thumbnail((1600, 1600))
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()


def iter_document(data: bytes, suffix: str):
    """返回带位置的片段；不运行宏，不提取压缩包中的任意文件到磁盘。"""
    stream = io.BytesIO(data)
    if suffix in {".txt", ".md", ".csv"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030")
        yield {"label": "正文", "text": text}
        return
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError("PDF 已加密")
        import pypdfium2 as pdfium
        document = None
        try:
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                yield {"label": f"第 {i + 1} 页", "text": text}
                # 无文字页面交给视觉模型；有文字页面的图片也参与理解。
                if not text.strip():
                    document = document or pdfium.PdfDocument(data)
                    pdf_page = document[i]
                    bitmap = pdf_page.render(scale=1.5)
                    try:
                        buf = io.BytesIO()
                        bitmap.to_pil().save(buf, "PNG")
                        yield {"label": f"第 {i + 1} 页扫描图", "image": image_url(buf.getvalue())}
                    finally:
                        bitmap.close()
                        pdf_page.close()
                else:
                    for n, img in enumerate(page.images):
                        yield {"label": f"第 {i + 1} 页图片 {n + 1}", "image": image_url(img.data)}
        finally:
            if document is not None:
                document.close()
        return
    if suffix not in {".docx", ".xlsx", ".pptx"}:
        raise ValueError("暂不支持此附件格式")
    with zipfile.ZipFile(stream) as archive:
        if sum(x.file_size for x in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("附件解压后超过 512 MB，请拆分文件")
    stream.seek(0)
    if suffix == ".docx":
        from docx import Document
        document = Document(stream)
        for i, para in enumerate(document.paragraphs):
            yield {"label": f"段落 {i + 1}", "text": para.text}
        for i, table in enumerate(document.tables):
            yield {"label": f"表格 {i + 1}", "text": "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)}
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            for sheet in book:
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    yield {"label": f"工作表 {sheet.title} 第 {i + 1} 行", "text": " | ".join(str(v) if v is not None else "" for v in row)}
        finally:
            book.close()
    else:
        from pptx import Presentation
        deck = Presentation(stream)
        for i, slide in enumerate(deck.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    yield {"label": f"幻灯片 {i + 1}", "text": shape.text}
                if shape.has_table:
                    yield {"label": f"幻灯片 {i + 1} 表格", "text": "\n".join(" | ".join(c.text for c in r.cells) for r in shape.table.rows)}
                if hasattr(shape, "image"):
                    yield {"label": f"幻灯片 {i + 1} 图片", "image": image_url(shape.image.blob)}
    if suffix != ".pptx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in archive.namelist():
                if "/media/" in name and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
                    yield {"label": f"嵌入图片 {Path(name).name}", "image": image_url(archive.read(name))}
    return


def parse_document(data: bytes, suffix: str):
    return list(iter_document(data, suffix))


@observed('media.resolve_media')
def resolve_media(account, message, max_mb):
    from ..chat_helpers import _resolve_account_dir
    from ..media_helpers import _resolve_media_path_for_kind, _read_and_maybe_decrypt_media, _fallback_search_media_by_file_id, _resolve_account_wxid_dir
    account_dir = _resolve_account_dir(account)
    kind, raw = message["kind"], message["media"]
    md5 = raw.get("imageMd5" if kind == "image" else "fileMd5", "")
    path = _resolve_media_path_for_kind(account_dir, kind=kind, md5=md5, username=message["username"], allow_fallback_scan=False)
    if path is None and kind == "image" and raw.get("imageFileId"):
        path = _fallback_search_media_by_file_id(str(_resolve_account_wxid_dir(account_dir) or ''), raw["imageFileId"], kind="image", username=message["username"], allow_global_scan=False)
    if path is None:
        raise ValueError("本机附件或图片缺失，请先在微信中下载并刷新")
    path = Path(path)
    if path.stat().st_size > max_mb * 1024 * 1024:
        raise ValueError(f"附件超过 {max_mb} MB，请调整上限后重试")
    if kind == "image":
        data, _ = _read_and_maybe_decrypt_media(path, account_dir=account_dir)
        return data, ".image"
    return path.read_bytes(), Path(raw.get("title") or path.name).suffix.lower()


class MediaService:
    def __init__(self, store, models):
        self.store, self.models = store, models

    async def enrich(self, account, message, options, vision_profile, checkpoint, unit_callback=None):
        if message["kind"] not in {"image", "file"}:
            return message
        return await self._enrich(account, message, options, vision_profile, checkpoint, unit_callback)

    @observed('media.enrich')
    async def _enrich(self, account, message, options, vision_profile, checkpoint, unit_callback=None):
        result = dict(message)
        if not options.get("media", True):
            return result | {"coverage": "未启用媒体分析"}
        try:
            data, suffix = await asyncio.to_thread(resolve_media, account, message, options.get("max_attachment_mb", 20))
            diagnostic_event('media.located', suffix=suffix, bytes=len(data))
            signature = {"version": 1, "profile": vision_profile.get("id"), "revision": vision_profile.get("revision"), "suffix": suffix}
            key = hashlib.sha256(account.encode() + data + json.dumps(signature, sort_keys=True).encode()).hexdigest()
            cached = self.store.get("media_cache", key)
            if cached:
                diagnostic_event('media.cache', cached=True, count=cached.get('units', 1))
                if unit_callback:
                    for n in range(cached.get('units', 1)):
                        unit_callback(f'缓存图片 {n + 1}', cached=True)
                return result | {"text": result["text"] + "\n" + cached["text"], "coverage": "已分析（缓存）"}
            parts = iter([{ "label": "图片", "image": image_url(data)}]) if suffix == ".image" else iter_document(data, suffix)
            text, index, units = [], -1, 0
            local_text = []
            while True:
                checkpoint()
                unit_started = time.monotonic()
                diagnostic_event('media.page.started', index=index+1)
                part = await asyncio.to_thread(next, parts, None)
                if part is None:
                    diagnostic_event('media.page.finished', index=index+1, complete=True)
                    break
                index += 1
                checkpoint()
                label = part["label"]
                if "image" in part:
                    if not vision_profile.get("vision"):
                        raise ValueError("该附件包含图片，请配置视觉模型后重试")
                    partial_key = f"{key}:{index}"
                    partial = self.store.get("media_cache", partial_key)
                    units += 1
                    if unit_callback:
                        unit_callback(label, cached=bool(partial))
                    if partial:
                        diagnostic_event('media.page.cache', index=index, cached=True)
                        description = partial["text"]
                    else:
                        description = await self.models.invoke(vision_profile, f"描述这份聊天资料中的{label}，完整提取可辨认文字、表格和关键信息。不能辨认的内容请说明。", images=[part["image"]], account=account)
                        self.store.put("media_cache", {"text": description}, id=partial_key, account=account)
                    text.append(f"[{label}] {description}")
                else:
                    text.append(f"[{label}] {part.get('text', '')}")
                    local_text.append(text[-1])
                    # 仅记录本地提取文字，独立于视觉模型输出，供离线检索复用。
                    local_id = hashlib.sha256(f"{account}:{message['username']}:{message['anchor']}".encode()).hexdigest()
                    self.store.put('local_media_text', {'text': '\n'.join(local_text), 'username': message['username'],
                        'anchor': message['anchor'], 'file_hash': hashlib.sha256(data).hexdigest(), 'version': 1}, id=local_id, account=account)
                diagnostic_event('media.page.finished', index=index, duration_ms=(time.monotonic()-unit_started)*1000)
            combined = "\n".join(text)
            self.store.put("media_cache", {"text": combined, "units": units}, id=key, account=account)
            return result | {"text": result["text"] + "\n" + combined, "coverage": "已分析"}
        except (ValueError, OSError, zipfile.BadZipFile) as exc:
            diagnostic_event('media.failed', level=logging.WARNING, error=exc, reason_code=media_failure_reason(exc))
            return result | {"coverage": str(exc), "text": result["text"] + "\n[附件未分析]"}


def media_failure_reason(error):
    # 仅将本地错误映射到固定分类，不记录异常正文。
    text = str(error)
    for needle, code in [('大小', 'size_limit'), ('超过', 'size_limit'), ('加密', 'encrypted'), ('损坏', 'corrupt'), ('不存在', 'missing'), ('找到', 'missing'), ('视觉', 'vision_required'), ('支持', 'unsupported')]:
        if needle in text:
            return code
    return 'parse_failed'
